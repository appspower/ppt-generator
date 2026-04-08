"""표 슬라이드 렌더러."""

from pptx.presentation import Presentation
from pptx.slide import Slide
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from .base import BaseRenderer
from ..styles import (
    CONTENT_LEFT, CONTENT_TOP, CONTENT_WIDTH, CONTENT_HEIGHT,
    COLOR_TABLE_HEADER, COLOR_TABLE_ROW_ALT, COLOR_TEXT,
    FONT_BODY, FONT_SIZE_TABLE_HEADER, FONT_SIZE_TABLE_BODY,
)
from ...models.schema import TableSlide


class TableRenderer(BaseRenderer):

    def render(self, prs: Presentation, slide_def: TableSlide) -> Slide:
        slide = self.add_blank_slide(prs)
        self.add_title(slide, slide_def.title)

        rows = len(slide_def.rows) + 1  # +1 for header
        cols = len(slide_def.headers)

        table_height = min(Inches(4.5), Inches(0.4) * rows)
        table = slide.shapes.add_table(
            rows, cols,
            CONTENT_LEFT, CONTENT_TOP + Inches(0.2),
            CONTENT_WIDTH, table_height,
        ).table

        # Column widths - 균등 분배
        col_width = int(CONTENT_WIDTH / cols)
        for i in range(cols):
            table.columns[i].width = col_width

        # Header row
        for j, header in enumerate(slide_def.headers):
            cell = table.cell(0, j)
            cell.text = header
            self._style_cell(cell, is_header=True)

        # Data rows
        for i, row_data in enumerate(slide_def.rows):
            for j, value in enumerate(row_data):
                cell = table.cell(i + 1, j)
                cell.text = str(value)
                self._style_cell(cell, is_header=False, is_alt_row=(i % 2 == 1))

        self.add_footnote(slide, slide_def.footnote)
        return slide

    def _style_cell(self, cell, is_header: bool, is_alt_row: bool = False):
        """셀 스타일 적용."""
        from pptx.oxml.ns import qn

        # 배경색
        if is_header:
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLOR_TABLE_HEADER
        elif is_alt_row:
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLOR_TABLE_ROW_ALT
        else:
            cell.fill.background()

        # 텍스트 스타일
        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.name = FONT_BODY
            if is_header:
                paragraph.font.size = FONT_SIZE_TABLE_HEADER
                paragraph.font.bold = True
                paragraph.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            else:
                paragraph.font.size = FONT_SIZE_TABLE_BODY
                paragraph.font.color.rgb = COLOR_TEXT

        # 여백
        cell.margin_left = Inches(0.1)
        cell.margin_right = Inches(0.1)
        cell.margin_top = Inches(0.05)
        cell.margin_bottom = Inches(0.05)
